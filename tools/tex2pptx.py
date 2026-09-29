#!/usr/bin/env python3
"""Convert the Beamer lecture decks to unstyled PowerPoint files.

The decks in ``lecture_*/slides/lecture_*.tex`` are written for Beamer's Madrid
theme.  We want the *content* in Google Slides with Google's own default
("blank" presentation) look, so this script re-renders each frame into a copy of
``Untitled presentation.pptx`` -- a blank Google Slides deck exported as .pptx.
That file supplies the theme, master and layouts; we only add slides.

Images are deliberately dropped (see README of the task): the originals live in
the Beamer PDFs and will be replaced by hand.

Every slide except the title page gets the yellow "UNEDITED!" starburst copied
verbatim from the sample deck, as a bookkeeping marker to delete after editing.
"""

from __future__ import annotations

import copy
import math
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"a": A, "p": P, "r": R}


def q(tag: str) -> str:
    prefix, local = tag.split(":")
    return "{%s}%s" % (NS[prefix], local)


# --------------------------------------------------------------------------
# geometry (EMU), matching the sample deck's own placeholder positions
# --------------------------------------------------------------------------
SLIDE_W, SLIDE_H = 9144000, 5143500
BODY_L, BODY_W = 311700, 8520600
TITLE_T, TITLE_H = 445025, 572700
BODY_T, BODY_H = 1152475, 3858000
FULL_T, FULL_H = 445025, 4565450  # untitled frames use the whole canvas

EMU_PER_PT = 12700
BODY_W_PT = BODY_W / EMU_PER_PT          # ~671 pt
BODY_H_PT = BODY_H / EMU_PER_PT          # ~304 pt
FULL_H_PT = FULL_H / EMU_PER_PT          # ~359 pt
TEXT_INSET_PT = 2 * 91425 / EMU_PER_PT   # placeholder l/r insets

MAX_BODY_PT, MIN_BODY_PT = 18, 9
MAX_TITLE_PT, MIN_TITLE_PT = 28, 14

# Arial is roughly half an em per character for mixed-case prose; Courier is
# a true monospace at 0.6 em.
PROP_CHAR_EM = 0.50
MONO_CHAR_EM = 0.60
LINE_SPACING = 1.20
# Greedy word wrap needs more lines than a pure character count suggests; the
# penalty is worse in narrow table cells than in a full-width paragraph.
WRAP_SLACK = 1.20
FLOW_GAP_PT = 8.0      # breathing room between stacked shapes
BOX_INSET_PT = 7.2     # a text box's own top + bottom insets

BULLETS = ["●", "○", "■", "●", "○"]
LVL0_MARL, LVL0_IND = 457200, -342900
LVLN_MARL, LVLN_IND = 457200, -317500


# --------------------------------------------------------------------------
# LaTeX -> inline runs
# --------------------------------------------------------------------------
SYMBOLS = {
    # Only codepoints that Arial actually has -- Google Slides' default font.
    r"\rightarrow": "\u2192", r"\longrightarrow": "\u2192",
    r"\Rightarrow": "\u2192", r"\leftarrow": "\u2190",
    r"\approx": "\u2248", r"\times": "\u00d7", r"\cdot": "\u00b7",
    r"\geq": "\u2265", r"\leq": "\u2264", r"\neq": "\u2260",
    r"\gg": ">>", r"\ll": "<<", r"\pm": "\u00b1",
    r"\ldots": "\u2026", r"\dots": "\u2026", r"\cdots": "\u2026",
    r"\to": "\u2192", r"\bullet": "\u2022", r"\checkmark": "\u221a",
    r"\alpha": "\u03b1", r"\beta": "\u03b2", r"\Delta": "\u0394",
    r"\infty": "\u221e", r"\sim": "~", r"\%": "%", r"\$": "$",
    r"\&": "&", r"\_": "_", r"\#": "#", r"\{": "{", r"\}": "}",
    r"\textbackslash": "\\", r"\today": "", r"\,": "\u2009", r"\;": " ",
    r"\:": " ", r"\!": "", r"\quad": "  ", r"\qquad": "    ",
}

# commands whose single argument is simply unwrapped
UNWRAP_1 = {
    "text", "mathrm", "mbox", "centering", "textnormal", "mathbf",
    "textsc", "underline", "uline", "textsf", "textrm",
}
# font-size / layout switches that carry no argument
DROP_0 = {
    "small", "footnotesize", "scriptsize", "tiny", "normalsize", "large",
    "Large", "LARGE", "huge", "Huge", "centering", "raggedright",
    "raggedleft", "bfseries", "itshape", "ttfamily", "noindent", "par",
    "hline", "toprule", "midrule", "bottomrule", "pause", "medskip",
    "bigskip", "smallskip", "clearpage", "newpage", "left", "right",
    "displaystyle", "arraystretch", "columnwidth", "textwidth", "linewidth",
}
# commands whose arguments are all discarded (name -> number of {} args)
DROP_N = {
    "vspace": 1, "hspace": 1, "includegraphics": 1, "label": 1, "ref": 1,
    "cite": 1, "usepackage": 1, "documentclass": 1, "vfill": 0, "hfill": 0,
    "renewcommand": 2, "newcommand": 2, "setlength": 2, "addtolength": 2,
    "color": 1, "textcolor": 2, "colorbox": 2, "def": 2, "usetheme": 1,
    "usecolortheme": 1, "usetikzlibrary": 1, "graphicspath": 1,
}


VERBATIM_RE = re.compile(
    r"\\begin\{(verbatim|Verbatim|lstlisting)\}.*?\\end\{\1\}", re.S)


def strip_comments(tex: str) -> str:
    """Drop % comments, but never inside a verbatim body (see the 20% in L4)."""
    held = []

    def hold(m):
        held.append(m.group(0))
        return "\x00VERB%d\x00" % (len(held) - 1)

    tex = VERBATIM_RE.sub(hold, tex)
    out = []
    for line in tex.split("\n"):
        i, n = 0, len(line)
        while i < n:
            if line[i] == "%" and (i == 0 or line[i - 1] != "\\"):
                line = line[:i]
                break
            i += 1
        out.append(line)
    tex = "\n".join(out)
    return re.sub(r"\x00VERB(\d+)\x00", lambda m: held[int(m.group(1))], tex)


def match_brace(s: str, i: int, open_c="{", close_c="}") -> int:
    """Index just past the group starting at s[i] == open_c."""
    depth = 0
    while i < len(s):
        c = s[i]
        if c == "\\":
            i += 2
            continue
        if c == open_c:
            depth += 1
        elif c == close_c:
            depth -= 1
            if depth == 0:
                return i + 1
        i += 1
    return len(s)


def read_arg(s: str, i: int):
    """Read one {...} argument starting at or after i; return (content, next_i)."""
    while i < len(s) and s[i] in " \n\t":
        i += 1
    if i < len(s) and s[i] == "{":
        j = match_brace(s, i)
        return s[i + 1:j - 1], j
    if i < len(s) and s[i] == "\\":  # e.g. \fbox\textbf -- unusual, bail out
        return "", i
    if i < len(s):
        return s[i], i + 1
    return "", i


def read_opt(s: str, i: int):
    """Read one optional [...] argument; return (content or None, next_i)."""
    j = i
    while j < len(s) and s[j] in " \n\t":
        j += 1
    if j < len(s) and s[j] == "[":
        k = match_brace(s, j, "[", "]")
        return s[j + 1:k - 1], k
    return None, i


class Run(dict):
    pass


def mkrun(text, style):
    r = Run(text=text)
    r.update(style)
    return r


def parse_inline(s: str, style: dict | None = None) -> list:
    """Turn a LaTeX fragment into a flat list of styled runs."""
    style = style or {"b": False, "i": False, "mono": False, "link": None}
    runs, buf, i, n = [], [], 0, len(s)

    def flush():
        if buf:
            runs.append(mkrun("".join(buf), style))
            buf.clear()

    def sub(inner, **over):
        flush()
        st = dict(style)
        st.update(over)
        runs.extend(parse_inline(inner, st))

    while i < n:
        c = s[i]

        if c == "\\":
            m = re.match(r"\\([a-zA-Z]+)\*?", s[i:])
            if not m:  # escaped punctuation or a line break
                if s[i:i + 2] == "\\\\":
                    _, j = read_opt(s, i + 2)
                    buf.append("\n")
                    i = j
                    continue
                if s[i:i + 2] in ("\\[", "\\]"):   # display math delimiters
                    buf.append("\n")
                    i += 2
                    continue
                if s[i:i + 2] in ("\\(", "\\)"):   # inline math delimiters
                    i += 2
                    continue
                if s[i:i + 2] in SYMBOLS:
                    buf.append(SYMBOLS[s[i:i + 2]])
                    i += 2
                    continue
                i += 1
                continue

            cmd, j = m.group(1), i + m.end()

            if "\\" + cmd in SYMBOLS:
                buf.append(SYMBOLS["\\" + cmd])
                i = j
                continue
            if cmd in ("textbf", "bf"):
                arg, i = read_arg(s, j)
                sub(arg, b=True)
                continue
            if cmd in ("textit", "emph", "it", "textsl"):
                arg, i = read_arg(s, j)
                sub(arg, i=True)
                continue
            if cmd in ("texttt", "tt", "path"):
                arg, i = read_arg(s, j)
                sub(arg, mono=True)
                continue
            if cmd == "verb":
                delim = s[j]
                k = s.index(delim, j + 1)
                sub(s[j + 1:k], mono=True)
                i = k + 1
                continue
            if cmd == "url":
                arg, i = read_arg(s, j)
                flush()
                st = dict(style)
                st["link"] = arg
                runs.append(mkrun(arg, st))
                continue
            if cmd == "href":
                url, j = read_arg(s, j)
                txt, i = read_arg(s, j)
                flush()
                st = dict(style)
                st["link"] = url
                runs.extend(parse_inline(txt, st))
                continue
            if cmd in ("fbox", "framebox", "boxed"):
                arg, i = read_arg(s, j)
                flush()
                buf.append("[ ")
                sub(arg)
                buf.append(" ]")
                continue
            if cmd in ("textcolor", "colorbox"):
                _, j = read_arg(s, j)
                arg, i = read_arg(s, j)
                sub(arg)
                continue
            if cmd == "fcolorbox":
                _, j = read_arg(s, j)
                _, j = read_arg(s, j)
                arg, i = read_arg(s, j)
                sub(arg)
                continue
            if cmd == "frac":
                num, j = read_arg(s, j)
                den, i = read_arg(s, j)
                sub(num)
                buf.append("/")
                sub(den)
                continue
            if cmd == "IfFileExists":
                for _ in range(3):
                    _, j = read_arg(s, j)
                i = j
                continue
            if cmd in DROP_N:
                _, j = read_opt(s, j)
                for _ in range(DROP_N[cmd]):
                    _, j = read_arg(s, j)
                i = j
                continue
            if cmd in UNWRAP_1:
                arg, i = read_arg(s, j)
                sub(arg)
                continue
            if cmd in DROP_0:
                i = j
                continue
            # unknown command: keep any braced argument's text
            arg, i2 = read_arg(s, j)
            if i2 > j and arg:
                sub(arg)
                i = i2
            else:
                i = j
            continue

        if c == "{":
            k = match_brace(s, i)
            sub(s[i + 1:k - 1])
            i = k
            continue
        if c == "}":
            i += 1
            continue
        if c == "$":
            k = s.find("$", i + 1)
            k = n if k < 0 else k
            sub(s[i + 1:k])
            i = k + 1
            continue
        if s[i:i + 2] in ("``", "''"):
            buf.append("\u201c" if s[i:i + 2] == "``" else "\u201d")
            i += 2
            continue
        if s[i:i + 3] == "---":
            buf.append("\u2014")
            i += 3
            continue
        if s[i:i + 2] == "--":
            buf.append("\u2013")
            i += 2
            continue
        if c == "~":
            buf.append(" ")
            i += 1
            continue
        if c in "_^":
            # T_{in} -> T_in ; 10^6 stays as-is
            nxt, k = read_arg(s, i + 1)
            buf.append(c)
            sub(nxt)
            i = k
            continue
        buf.append(c)
        i += 1

    flush()
    # merge adjacent runs with identical styling and collapse whitespace runs
    merged = []
    for r in runs:
        if merged and all(merged[-1][k] == r[k] for k in ("b", "i", "mono", "link")):
            merged[-1]["text"] += r["text"]
        else:
            merged.append(r)
    for r in merged:
        r["text"] = re.sub(r"[ \t]*\n[ \t\n]*", "\n", r["text"])
    return merged


def clean_runs(runs: list) -> list:
    """Collapse runs of whitespace, drop empties, trim the ends."""
    out = []
    for r in runs:
        t = r["text"]
        if not r["mono"]:
            t = re.sub(r"[ \t]+", " ", t)
        if t:
            r = Run(r)
            r["text"] = t
            out.append(r)
    while out and not out[0]["text"].strip():
        out.pop(0)
    while out and not out[-1]["text"].strip():
        out.pop()
    if out:
        out[0] = Run(out[0], text=out[0]["text"].lstrip())
        out[-1] = Run(out[-1], text=out[-1]["text"].rstrip())
    return [r for r in out if r["text"]]


def runs_len(runs) -> int:
    return sum(len(r["text"]) for r in runs)


# --------------------------------------------------------------------------
# LaTeX -> block structure
# --------------------------------------------------------------------------
ENV_RE = re.compile(r"\\begin\{([A-Za-z*]+)\}")


def find_env(s: str, start: int = 0):
    """Locate the next top-level environment; returns (name, b0, b1, body) or None."""
    m = ENV_RE.search(s, start)
    if not m:
        return None
    name = m.group(1)
    i = m.end()
    # skip [opt] and {arg} directly after \begin{env}
    depth, j = 1, i
    open_re = re.compile(r"\\begin\{%s\}" % re.escape(name))
    close_re = re.compile(r"\\end\{%s\}" % re.escape(name))
    while j < len(s):
        mo = open_re.search(s, j)
        mc = close_re.search(s, j)
        if not mc:
            return name, m.start(), len(s), s[i:]
        if mo and mo.start() < mc.start():
            depth += 1
            j = mo.end()
        else:
            depth -= 1
            if depth == 0:
                return name, m.start(), mc.end(), s[i:mc.start()]
            j = mc.end()
    return name, m.start(), len(s), s[i:]


def split_items(body: str) -> list:
    """Split an itemize/enumerate body on top-level \\item."""
    parts, depth, i, last, n = [], 0, 0, 0, len(body)
    while i < n:
        if body.startswith("\\begin{", i):
            depth += 1
            i += 7
            continue
        if body.startswith("\\end{", i):
            depth -= 1
            i += 5
            continue
        if depth == 0 and body.startswith("\\item", i) and not re.match(r"\\item[a-zA-Z]", body[i:]):
            parts.append(body[last:i])
            last = i + 5
            i += 5
            continue
        i += 1
    parts.append(body[last:])
    return parts[1:] if parts and not parts[0].strip() else parts


def parse_blocks(s: str, depth: int = 0, align: str = "l") -> list:
    """Turn a frame body into an ordered list of block dicts."""
    blocks, pos = [], 0
    while True:
        found = find_env(s, pos)
        if not found:
            blocks += text_blocks(s[pos:], depth, align)
            break
        name, b0, b1, body = found
        blocks += text_blocks(s[pos:b0], depth, align)
        blocks += env_blocks(name, body, depth, align)
        pos = b1
    return [b for b in blocks if b]


# environment name -> (number of [opt] args, number of {req} args) to discard
ENV_ARGS = {
    "columns": (1, 0), "column": (0, 1), "minipage": (1, 1),
    "tabular": (1, 0), "tabularx": (1, 1), "adjustbox": (0, 1),
    "figure": (1, 0), "table": (1, 0), "itemize": (1, 0),
    "enumerate": (1, 0), "description": (1, 0), "tikzpicture": (1, 0),
    "lstlisting": (1, 0), "center": (0, 0),
}


def strip_env_args(name: str, body: str) -> str:
    n_opt, n_req = ENV_ARGS.get(name, (0, 0))
    i = 0
    for _ in range(n_opt):
        _, i = read_opt(body, i)
    for _ in range(n_req):
        _, i = read_arg(body, i)
    return body[i:]


def env_blocks(name: str, body: str, depth: int, align: str) -> list:
    if name not in ("verbatim", "Verbatim", "block", "exampleblock", "alertblock"):
        body = strip_env_args(name, body)
    if name in ("itemize", "enumerate", "description"):
        out = []
        for raw in split_items(body):
            out += item_blocks(raw, depth, "num" if name == "enumerate" else "bul")
        return out
    if name in ("center", "centering"):
        return parse_blocks(body, depth, "ctr")
    if name in ("columns", "column", "minipage", "figure", "table", "adjustbox",
                "small", "footnotesize", "scriptsize", "quotation", "flushleft"):
        return parse_blocks(body, depth, align)
    if name == "quote":
        return [dict(b, quote=True) for b in parse_blocks(body, depth + 1, align)]
    if name == "block":
        # \begin{block}{Title} -- the title is the first {...} of the body
        title, rest = "", body
        m = re.match(r"\s*\{", body)
        if m:
            j = match_brace(body, m.end() - 1)
            title, rest = body[m.end():j - 1], body[j:]
        out = []
        if title.strip():
            out.append({"kind": "para", "runs": clean_runs(parse_inline(title)),
                        "depth": depth, "align": align, "strong": True})
        return out + parse_blocks(rest, depth, align)
    if name in ("verbatim", "lstlisting", "Verbatim"):
        code = body.strip("\n")
        code = "\n".join(line.rstrip() for line in code.split("\n"))
        return [{"kind": "code", "text": code, "depth": depth}]
    if name == "tabular":
        return [parse_tabular(body, depth, align)]
    if name in ("tikzpicture", "pgfpicture", "overpic"):
        return []
    return parse_blocks(body, depth, align)


def item_blocks(raw: str, depth: int, style: str) -> list:
    """One \\item: its own text becomes a bullet, nested lists recurse deeper."""
    lead, rest = raw, ""
    found = find_env(raw)
    if found:
        lead, rest = raw[:found[1]], raw[found[1]:]
    runs = clean_runs(parse_inline(lead))
    out = []
    if runs:
        out.append({"kind": "bullet", "runs": runs, "depth": depth, "style": style,
                    "align": "l"})
    if rest.strip():
        out += parse_blocks(rest, depth + 1 if runs else depth, "l")
    return out


COLSPEC_P_RE = re.compile(r"p\{([0-9.]+)\\(?:text|line|column)width\}")


def colspec_weights(spec: str, ncols: int):
    """Relative column widths from a LaTeX column spec, or None if unspecified."""
    if not spec:
        return None
    fracs, i = [], 0
    while i < len(spec):
        m = COLSPEC_P_RE.match(spec, i)
        if m:
            fracs.append(float(m.group(1)))
            i = m.end()
            continue
        if spec[i] in "lcrX":
            fracs.append(None)
        i += 1
    if len(fracs) != ncols or all(f is None for f in fracs):
        return None
    known = [f for f in fracs if f is not None]
    fill = (max(0.0, 1.0 - sum(known)) / max(1, len(fracs) - len(known))) or 0.15
    return [f if f is not None else fill for f in fracs]


def parse_tabular(body: str, depth: int, align: str) -> dict:
    spec, k = read_arg(body, 0)
    body = body[k:]
    body = re.sub(r"\\hline|\\toprule|\\midrule|\\bottomrule", "", body)
    rows = []
    for raw in re.split(r"\\\\(?:\s*\[[^\]]*\])?", body):
        if not raw.strip():
            continue
        cells, depth_b, last, i = [], 0, 0, 0
        while i < len(raw):
            ch = raw[i]
            if ch == "\\":
                i += 2
                continue
            if ch == "{":
                depth_b += 1
            elif ch == "}":
                depth_b -= 1
            elif ch == "&" and depth_b == 0:
                cells.append(raw[last:i])
                last = i + 1
            i += 1
        cells.append(raw[last:])
        rows.append([clean_runs(parse_inline(c)) for c in cells])
    width = max((len(r) for r in rows), default=0)
    rows = [r + [[]] * (width - len(r)) for r in rows]
    return {"kind": "table", "rows": rows, "depth": depth, "align": align,
            "weights": colspec_weights(spec, width)}


DISPLAY_MATH_RE = re.compile(r"\\\[(.*?)\\\]", re.S)


def text_blocks(s: str, depth: int, align: str) -> list:
    if not s.strip():
        return []
    out = []
    # \[ ... \] display math becomes its own centered line
    pieces, pos = [], 0
    for m in DISPLAY_MATH_RE.finditer(s):
        pieces.append(("t", s[pos:m.start()]))
        pieces.append(("m", m.group(1)))
        pos = m.end()
    pieces.append(("t", s[pos:]))
    if len(pieces) > 1:
        for kind, chunk in pieces:
            if kind == "m":
                runs = clean_runs(parse_inline(chunk))
                if runs:
                    out.append({"kind": "para", "runs": runs, "depth": depth,
                                "align": "ctr", "math": True})
            else:
                out += text_blocks(chunk, depth, align)
        return out
    for chunk in re.split(r"\n[ \t]*\n", s):
        if not chunk.strip():
            continue
        runs = clean_runs(parse_inline(chunk))
        if runs:
            out.append({"kind": "para", "runs": runs, "depth": depth, "align": align})
    return out


# --------------------------------------------------------------------------
# document -> slide specs
# --------------------------------------------------------------------------
def parse_document(tex: str) -> tuple[dict, list]:
    tex = strip_comments(tex)
    meta = {}
    for key in ("title", "subtitle", "author", "date", "institute"):
        m = re.search(r"\\%s\s*\{" % key, tex)
        if m:
            j = match_brace(tex, m.end() - 1)
            meta[key] = "".join(r["text"] for r in clean_runs(parse_inline(tex[m.end():j - 1])))
    body = tex.split(r"\begin{document}", 1)[1].split(r"\end{document}")[0]

    # walk the document top level, picking up \section and frames in order
    slides, pos = [], 0
    frame_re = re.compile(r"\\begin\{frame\}")
    sect_re = re.compile(r"\\section(\*?)\s*\{")
    title_re = re.compile(r"\\frame\s*\{\s*\\titlepage\s*\}")
    while pos < len(body):
        mf = frame_re.search(body, pos)
        ms = sect_re.search(body, pos)
        mt = title_re.search(body, pos)
        cands = [m for m in (mf, ms, mt) if m]
        if not cands:
            break
        m = min(cands, key=lambda x: x.start())
        if m is mt:
            slides.append({"type": "title"})
            pos = m.end()
        elif m is ms:
            j = match_brace(body, m.end() - 1)
            name = "".join(r["text"] for r in clean_runs(parse_inline(body[m.end():j - 1])))
            starred = m.group(1) == "*"
            slides.append({"type": "section", "name": name, "starred": starred})
            pos = j
        else:
            found = find_env(body, m.start())
            _, b0, b1, fbody = found
            slides.append(parse_frame(fbody))
            pos = b1
    return meta, slides


def parse_frame(fbody: str) -> dict:
    i = 0
    _, i = read_opt(fbody, i)              # [fragile], [t], ...
    title = None
    while i < len(fbody) and fbody[i] in " \n\t":
        i += 1
    if i < len(fbody) and fbody[i] == "{":
        j = match_brace(fbody, i)
        title = fbody[i + 1:j - 1]
        i = j
    rest = fbody[i:]
    m = re.search(r"\\frametitle\s*\{", rest)
    if title is None and m:
        j = match_brace(rest, m.end() - 1)
        title = rest[m.end():j - 1]
        rest = rest[:m.start()] + rest[j:]
    toc = bool(re.search(r"\\tableofcontents", rest))
    return {"type": "frame",
            "title": clean_runs(parse_inline(title)) if title else None,
            "toc": toc,
            "blocks": [] if toc else parse_blocks(rest)}


def expand(meta: dict, slides: list) -> list:
    """Replace \\tableofcontents frames with a real outline and add dividers."""
    sections = [s["name"] for s in slides if s["type"] == "section" and not s["starred"]]
    out = []
    for s in slides:
        if s["type"] == "section":
            if not s["starred"]:
                out.append({"type": "section", "name": s["name"]})
            continue
        if s.get("toc"):
            if any(o.get("type") == "outline" for o in out):
                continue  # per-section outlines become section header slides
            out.append({"type": "outline", "sections": sections,
                        "title": s["title"]})
            continue
        out.append(s)
    return out


# --------------------------------------------------------------------------
# rendering helpers
# --------------------------------------------------------------------------
def set_rpr(rpr, size_pt, run):
    rpr.set("lang", "en")
    if run.get("b"):
        rpr.set("b", "1")
    if run.get("i"):
        rpr.set("i", "1")
    rpr.set("sz", str(int(round(size_pt * 100))))
    if run.get("mono"):
        for tag in ("latin", "ea", "cs", "sym"):
            el = etree.SubElement(rpr, q("a:" + tag))
            el.set("typeface", "Courier New")


def add_paragraph(txbody, runs, *, size_pt, level=0, bullet=None,
                  align="l", space_before_pt=0.0, part=None):
    p = etree.SubElement(txbody, q("a:p"))
    ppr = etree.SubElement(p, q("a:pPr"))
    ppr.set("lvl", str(min(level, 8)))
    ppr.set("rtl", "0")
    ppr.set("algn", align)
    if bullet:
        marl = (LVL0_MARL if level == 0 else LVL0_MARL) * (level + 1)
        ind = LVL0_IND if level == 0 else LVLN_IND
        ppr.set("marL", str(marl))
        ppr.set("indent", str(ind))
    else:
        ppr.set("marL", str(LVL0_MARL * level))
        ppr.set("indent", "0")
    sb = etree.SubElement(ppr, q("a:spcBef"))
    etree.SubElement(sb, q("a:spcPts")).set("val", str(int(space_before_pt * 100)))
    sa = etree.SubElement(ppr, q("a:spcAft"))
    etree.SubElement(sa, q("a:spcPts")).set("val", "0")
    if bullet == "num":
        etree.SubElement(ppr, q("a:buSzPts")).set("val", str(int(size_pt * 100)))
        bu = etree.SubElement(ppr, q("a:buAutoNum"))
        bu.set("type", "arabicPeriod")
    elif bullet == "bul":
        etree.SubElement(ppr, q("a:buSzPts")).set("val", str(int(size_pt * 100)))
        etree.SubElement(ppr, q("a:buChar")).set("char", BULLETS[min(level, len(BULLETS) - 1)])
    else:
        etree.SubElement(ppr, q("a:buNone"))

    for run in runs:
        pieces = run["text"].split("\n")
        for k, piece in enumerate(pieces):
            if k:
                etree.SubElement(p, q("a:br"))
            if not piece:
                continue
            r = etree.SubElement(p, q("a:r"))
            rpr = etree.SubElement(r, q("a:rPr"))
            set_rpr(rpr, size_pt, run)
            if run.get("link") and part is not None:
                rid = part.relate_to(run["link"],
                                     "http://schemas.openxmlformats.org/officeDocument/"
                                     "2006/relationships/hyperlink", is_external=True)
                hl = etree.SubElement(rpr, q("a:hlinkClick"))
                hl.set(q("r:id"), rid)
            t = etree.SubElement(r, q("a:t"))
            t.text = piece
    epr = etree.SubElement(p, q("a:endParaRPr"))
    epr.set("sz", str(int(round(size_pt * 100))))
    return p


def style_cell(tc):
    """Thin grey borders, no fill -- Google Slides' own default table look."""
    tcPr = tc._tc.find(q("a:tcPr"))
    if tcPr is None:
        tcPr = etree.SubElement(tc._tc, q("a:tcPr"))
    for name in ("lnL", "lnR", "lnT", "lnB"):
        ln = etree.SubElement(tcPr, q("a:" + name))
        ln.set("w", "9525")
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        fill = etree.SubElement(ln, q("a:solidFill"))
        etree.SubElement(fill, q("a:srgbClr")).set("val", "999999")
        etree.SubElement(ln, q("a:prstDash")).set("val", "solid")
    etree.SubElement(tcPr, q("a:noFill"))
    tcPr.set("marL", "68580")
    tcPr.set("marR", "68580")
    tcPr.set("marT", "34290")
    tcPr.set("marB", "34290")


def para_height_pt(nchars, size_pt, level, bullet, width_pt, extra_lines=0,
                   space_before_pt=0.0, mono=False):
    indent = 36.0 * (level + (1 if bullet else 0))
    avail = max(60.0, width_pt - TEXT_INSET_PT - indent)
    em = MONO_CHAR_EM if mono else PROP_CHAR_EM
    cpl = max(8.0, avail / (em * size_pt))
    lines = max(1, math.ceil(nchars / cpl)) + extra_lines
    return lines * size_pt * LINE_SPACING + space_before_pt


def flow_height(blocks, size_pt, width_pt):
    """Total height of a run of text blocks at the given base font size."""
    total = 0.0
    for b in blocks:
        total += block_height(b, size_pt, width_pt)
    return total


def lvl_size(base_pt, level):
    return max(MIN_BODY_PT - 1, base_pt - min(level, 2))


def block_height(b, size_pt, width_pt):
    if b["kind"] == "table":
        return table_height(b, size_pt, width_pt)
    if b["kind"] == "code":
        cs = code_size(b["text"], size_pt, width_pt)
        nlines = b["text"].count("\n") + 1
        return nlines * cs * LINE_SPACING + 0.5 * size_pt
    lvl = b["depth"]
    s = lvl_size(size_pt, lvl)
    bullet = b["kind"] == "bullet"
    nl = sum(r["text"].count("\n") for r in b["runs"])
    sb = 0.45 * s if bullet else 0.6 * s
    return para_height_pt(runs_len(b["runs"]), s, lvl, bullet, width_pt,
                          extra_lines=nl, space_before_pt=sb)


def code_size(text, base_pt, width_pt):
    longest = max((len(l) for l in text.split("\n")), default=1)
    fit = (width_pt - TEXT_INSET_PT - 20) / (MONO_CHAR_EM * max(longest, 1))
    return max(7.0, min(base_pt - 3, fit))


def table_weights(b):
    """Normalised column widths: the LaTeX spec if given, else by content."""
    ncols = max(1, len(b["rows"][0]))
    w = b.get("weights")
    if not w:
        longest = [max((runs_len(r[c]) for r in b["rows"]), default=1) or 1
                   for c in range(ncols)]
        total = sum(longest)
        w = [max(0.12, min(0.65, x / total)) for x in longest]
    total = sum(w) or 1.0
    return [x / total for x in w]


def table_height(b, size_pt, width_pt):
    s = max(MIN_BODY_PT, size_pt - 2)
    weights = table_weights(b)
    h = 0.5 * size_pt + 0.4 * size_pt  # top gap + a little slack
    for row in b["rows"]:
        rows_lines = 1
        for ci, cell in enumerate(row):
            colw = width_pt * weights[ci]
            cpl = max(4.0, (colw - 14) / (PROP_CHAR_EM * s * WRAP_SLACK))
            rows_lines = max(rows_lines, math.ceil(max(1, runs_len(cell)) / cpl))
        h += rows_lines * s * LINE_SPACING + 0.6 * s
    return h


def fit_size(blocks, avail_pt, width_pt):
    for size in range(MAX_BODY_PT, MIN_BODY_PT - 1, -1):
        if flow_height(blocks, float(size), width_pt) <= avail_pt:
            return float(size)
    return float(MIN_BODY_PT)


def fit_title(runs):
    size = float(MAX_TITLE_PT)
    n = runs_len(runs)
    while size > MIN_TITLE_PT and n * PROP_CHAR_EM * size > (BODY_W_PT - TEXT_INSET_PT):
        size -= 1
    return size


# --------------------------------------------------------------------------
# slide construction
# --------------------------------------------------------------------------
class Builder:
    def __init__(self, template: Path):
        self.prs = Presentation(str(template))
        self.layouts = {l.name: l for l in self.prs.slide_layouts}
        self.stamp = self._grab_stamp()
        self._clear()

    def _grab_stamp(self):
        for slide in self.prs.slides:
            for sh in slide.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip() == "UNEDITED!":
                    return copy.deepcopy(sh._element)
        raise RuntimeError("no UNEDITED! shape found in the template")

    def _clear(self):
        xml_slides = self.prs.slides._sldIdLst
        for sld in list(xml_slides):
            rId = sld.get(q("r:id"))
            self.prs.part.drop_rel(rId)
            xml_slides.remove(sld)

    # -- low-level shape helpers -------------------------------------------
    def _new(self, layout_name):
        return self.prs.slides.add_slide(self.layouts[layout_name])

    @staticmethod
    def _spTree(slide):
        return slide.shapes._spTree

    @staticmethod
    def _next_id(slide):
        ids = [int(e.get("id")) for e in slide.shapes._spTree.iter(q("p:cNvPr"))
               if e.get("id") and e.get("id").isdigit()]
        return max(ids + [1]) + 1

    def add_stamp(self, slide):
        el = copy.deepcopy(self.stamp)
        el.find(q("p:nvSpPr")).find(q("p:cNvPr")).set("id", str(self._next_id(slide)))
        self._spTree(slide).append(el)

    @staticmethod
    def placeholder(slide, idx):
        for ph in slide.placeholders:
            if ph.placeholder_format.idx == idx:
                return ph
        return None

    @staticmethod
    def drop(shape):
        shape._element.getparent().remove(shape._element)

    @staticmethod
    def place(shape, left, top, width, height):
        shape.left, shape.top = Emu(left), Emu(top)
        shape.width, shape.height = Emu(width), Emu(height)

    @staticmethod
    def clear_tf(shape):
        txbody = shape._element.find(q("p:txBody"))
        for p in txbody.findall(q("a:p")):
            txbody.remove(p)
        bodyPr = txbody.find(q("a:bodyPr"))
        for af in bodyPr.findall(q("a:normAutofit")) + bodyPr.findall(q("a:spAutoFit")):
            bodyPr.remove(af)
        etree.SubElement(bodyPr, q("a:noAutofit"))
        return txbody

    def textbox(self, slide, left, top, width, height, anchor="t"):
        sp = etree.SubElement(self._spTree(slide), q("p:sp"))
        nv = etree.SubElement(sp, q("p:nvSpPr"))
        cnv = etree.SubElement(nv, q("p:cNvPr"))
        cnv.set("id", str(self._next_id(slide)))
        cnv.set("name", "TextBox")
        etree.SubElement(nv, q("p:cNvSpPr")).set("txBox", "1")
        etree.SubElement(nv, q("p:nvPr"))
        sppr = etree.SubElement(sp, q("p:spPr"))
        xfrm = etree.SubElement(sppr, q("a:xfrm"))
        off = etree.SubElement(xfrm, q("a:off"))
        off.set("x", str(int(left)))
        off.set("y", str(int(top)))
        ext = etree.SubElement(xfrm, q("a:ext"))
        ext.set("cx", str(int(width)))
        ext.set("cy", str(int(height)))
        geom = etree.SubElement(sppr, q("a:prstGeom"))
        geom.set("prst", "rect")
        etree.SubElement(geom, q("a:avLst"))
        etree.SubElement(sppr, q("a:noFill"))
        tx = etree.SubElement(sp, q("p:txBody"))
        bp = etree.SubElement(tx, q("a:bodyPr"))
        bp.set("wrap", "square")
        bp.set("anchor", anchor)
        bp.set("lIns", "91425")
        bp.set("rIns", "91425")
        bp.set("tIns", "45700")
        bp.set("bIns", "45700")
        etree.SubElement(bp, q("a:noAutofit"))
        etree.SubElement(tx, q("a:lstStyle"))
        return tx

    # -- block emitters ----------------------------------------------------
    def emit_text_blocks(self, txbody, blocks, size_pt, part):
        for b in blocks:
            if b["kind"] == "code":
                cs = code_size(b["text"], size_pt, BODY_W_PT)
                first = True
                for line in b["text"].split("\n"):
                    add_paragraph(txbody, [Run(text=line or " ", b=False, i=False,
                                               mono=True, link=None)],
                                  size_pt=cs, level=max(b["depth"], 1), bullet=None,
                                  align="l",
                                  space_before_pt=(0.5 * size_pt if first else 0.0),
                                  part=part)
                    first = False
                continue
            lvl = b["depth"]
            s = lvl_size(size_pt, lvl)
            if b["kind"] == "bullet":
                add_paragraph(txbody, b["runs"], size_pt=s, level=lvl,
                              bullet=b["style"], align="l",
                              space_before_pt=0.45 * s, part=part)
            else:
                runs = b["runs"]
                if b.get("strong"):
                    runs = [Run(r, b=True) for r in runs]
                add_paragraph(txbody, runs, size_pt=s, level=lvl, bullet=None,
                              align=b.get("align", "l"),
                              space_before_pt=0.6 * s, part=part)

    def emit_table(self, slide, b, top, size_pt, width, left):
        rows, cols = len(b["rows"]), len(b["rows"][0])
        s = max(MIN_BODY_PT, size_pt - 2)
        height = int(table_height(b, size_pt, BODY_W_PT) * EMU_PER_PT)
        gf = slide.shapes.add_table(rows, cols, Emu(int(left)), Emu(int(top)),
                                    Emu(int(width)), Emu(height))
        tbl = gf.table
        # Drop PowerPoint's banded accent style: we want Google's plain look.
        tblPr = gf._element.graphic.graphicData.tbl.find(q("a:tblPr"))
        tblPr.set("firstRow", "0")
        tblPr.set("bandRow", "0")
        for sid in tblPr.findall(q("a:tableStyleId")):
            tblPr.remove(sid)
        for ci, wgt in enumerate(table_weights(b)):
            tbl.columns[ci].width = Emu(int(width * wgt))
        for ri, row in enumerate(b["rows"]):
            for ci, cell in enumerate(row):
                tc = tbl.cell(ri, ci)
                txbody = tc._tc.find(q("a:txBody"))
                for p in txbody.findall(q("a:p")):
                    txbody.remove(p)
                add_paragraph(txbody, cell or [Run(text="", b=False, i=False,
                                                   mono=False, link=None)],
                              size_pt=s, level=0, bullet=None, align="l",
                              space_before_pt=0.0)
                style_cell(tc)

    # -- slide kinds -------------------------------------------------------
    def title_slide(self, meta):
        slide = self._new("TITLE")
        t = self.placeholder(slide, 0)
        sub = self.placeholder(slide, 1)
        self.place(t, 311708, 744575, 8520600, 2052600)
        self.place(sub, 311700, 2834125, 8520600, 792600)
        tx = self.clear_tf(t)
        add_paragraph(tx, [Run(text=meta.get("title", ""), b=False, i=False,
                               mono=False, link=None)],
                      size_pt=40.0, level=0, bullet=None, align="l")
        tx = self.clear_tf(sub)
        lines = [meta.get("subtitle", ""), meta.get("author", "")]
        first = True
        for line in [l for l in lines if l]:
            add_paragraph(tx, [Run(text=line, b=False, i=False, mono=False, link=None)],
                          size_pt=18.0, level=0, bullet=None, align="l",
                          space_before_pt=0.0 if first else 6.0)
            first = False

    def section_slide(self, name):
        slide = self._new("SECTION_HEADER")
        t = self.placeholder(slide, 0)
        self.place(t, 311700, 1728125, 8520600, 1687300)
        tx = self.clear_tf(t)
        runs = clean_runs(parse_inline(name))
        size = 40.0
        while size > 20 and runs_len(runs) * PROP_CHAR_EM * size > 2 * (BODY_W_PT - TEXT_INSET_PT):
            size -= 2
        add_paragraph(tx, runs, size_pt=size, level=0, bullet=None, align="l")
        self.add_stamp(slide)

    def outline_slide(self, spec):
        blocks = [{"kind": "bullet", "runs": clean_runs(parse_inline(name)),
                   "depth": 0, "style": "num", "align": "l"}
                  for name in spec["sections"]]
        self.content_slide(spec.get("title") or clean_runs(parse_inline("Outline")), blocks)

    def content_slide(self, title_runs, blocks):
        has_table = any(b["kind"] == "table" for b in blocks)
        if title_runs and not has_table:
            self._simple_slide(title_runs, blocks)
        else:
            self._flow_slide(title_runs, blocks)

    def _simple_slide(self, title_runs, blocks):
        slide = self._new("TITLE_AND_BODY")
        part = slide.part
        t = self.placeholder(slide, 0)
        body = self.placeholder(slide, 1)
        self.place(t, BODY_L, TITLE_T, BODY_W, TITLE_H)
        self.place(body, BODY_L, BODY_T, BODY_W, BODY_H)
        tx = self.clear_tf(t)
        add_paragraph(tx, title_runs, size_pt=fit_title(title_runs), level=0,
                      bullet=None, align="l", part=part)
        tx = self.clear_tf(body)
        if blocks:
            size = fit_size(blocks, BODY_H_PT - 6, BODY_W_PT)
            self.emit_text_blocks(tx, blocks, size, part)
        else:
            add_paragraph(tx, [], size_pt=18.0, level=0, bullet=None)
        self.add_stamp(slide)

    def _flow_slide(self, title_runs, blocks):
        """Slides with tables (or no title) get hand-placed shapes."""
        slide = self._new("TITLE_ONLY")
        part = slide.part
        t = self.placeholder(slide, 0)
        if title_runs:
            self.place(t, BODY_L, TITLE_T, BODY_W, TITLE_H)
            tx = self.clear_tf(t)
            add_paragraph(tx, title_runs, size_pt=fit_title(title_runs), level=0,
                          bullet=None, align="l", part=part)
            top, avail = BODY_T, BODY_H_PT
        else:
            self.drop(t)
            top, avail = FULL_T, FULL_H_PT

        # group consecutive non-table blocks into single text boxes
        groups, cur = [], []
        for b in blocks:
            if b["kind"] == "table":
                if cur:
                    groups.append(("text", cur))
                    cur = []
                groups.append(("table", b))
            else:
                cur.append(b)
        if cur:
            groups.append(("text", cur))

        def group_h(kind, payload, size):
            if kind == "table":
                return table_height(payload, size, BODY_W_PT)
            return flow_height(payload, size, BODY_W_PT) + BOX_INSET_PT

        def total(size):
            return (sum(group_h(k, pl, size) for k, pl in groups)
                    + FLOW_GAP_PT * max(0, len(groups) - 1))

        size = float(MIN_BODY_PT)
        for cand in range(MAX_BODY_PT, MIN_BODY_PT - 1, -1):
            if total(float(cand)) <= avail - 6:
                size = float(cand)
                break

        y = top
        for kind, payload in groups:
            if kind == "text":
                h = int(group_h(kind, payload, size) * EMU_PER_PT)
                tx = self.textbox(slide, BODY_L, y, BODY_W, max(h, 200000))
                self.emit_text_blocks(tx, payload, size, part)
            else:
                ncols = len(payload["rows"][0])
                width = BODY_W if ncols > 1 else int(BODY_W * 0.6)
                left = BODY_L + (BODY_W - width) // 2
                h = int(group_h(kind, payload, size) * EMU_PER_PT)
                self.emit_table(slide, payload, y, size, width, left)
            y += h + int(FLOW_GAP_PT * EMU_PER_PT)
        self.add_stamp(slide)

    def save(self, path: Path):
        self.prs.save(str(path))


def convert(tex_path: Path, template: Path, out_path: Path):
    meta, raw = parse_document(tex_path.read_text(encoding="utf-8"))
    slides = expand(meta, raw)
    b = Builder(template)
    for spec in slides:
        if spec["type"] == "title":
            b.title_slide(meta)
        elif spec["type"] == "section":
            b.section_slide(spec["name"])
        elif spec["type"] == "outline":
            b.outline_slide(spec)
        else:
            if not spec["title"] and not spec["blocks"]:
                continue
            b.content_slide(spec["title"], spec["blocks"])
    b.save(out_path)
    return len(b.prs.slides._sldIdLst)


def main(argv):
    root = Path(__file__).resolve().parent.parent
    template = root / "Untitled presentation.pptx"
    targets = argv[1:] or sorted(str(p) for p in root.glob("lecture_*/slides/lecture_*.tex"))
    for t in targets:
        tex = Path(t)
        out = tex.with_suffix(".pptx")
        n = convert(tex, template, out)
        print(f"{tex.relative_to(root)} -> {out.relative_to(root)}  ({n} slides)")


if __name__ == "__main__":
    main(sys.argv)
